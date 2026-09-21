from fastapi.testclient import TestClient
import fitz
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import load_workbook

from api.main import app


def test_live_frontend_contract() -> None:
    with TestClient(app) as client:
        health = client.get("/api/health")
        assert health.status_code == 200
        assert health.json()["cases"] > 0
        assert health.json()["mode"] in {"full", "deterministic_only"}
        assert health.json()["circuit_breaker"] in {"closed", "open", "half_open"}
        assert health.json()["components"]

        metrics = client.get("/api/metrics")
        assert metrics.status_code == 200
        assert metrics.json()["automatically_cleared"] == 154
        assert metrics.json()["human_reviews"] == 20
        assert metrics.json()["top_flagged_issues"]
        assert metrics.json()["avg_processing_ms"] is not None

        info = client.get("/api")
        assert info.status_code == 200
        assert info.json()["cases_loaded"] == 520

        second_page = client.get("/api/cases?limit=500&offset=500")
        assert second_page.status_code == 200
        assert len(second_page.json()["items"]) == 20
        assert "created_at" in second_page.json()["items"][0]
        assert "category_confidence" in second_page.json()["items"][0]
        assert "assigned_to" in second_page.json()["items"][0]
        assert "review_status" in second_page.json()["items"][0]

        excel = client.get("/api/cases/export.xlsx")
        assert excel.status_code == 200
        assert excel.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment" in excel.headers["content-disposition"]
        workbook = load_workbook(BytesIO(excel.content), read_only=True, data_only=False)
        assert workbook.sheetnames == [
            "Summary", "All Cases", "Field Comparisons", "Document Evidence",
            "Audit Trail", "Review Queue",
        ]
        assert workbook["All Cases"].max_row == 521
        assert workbook["All Cases"]["A2"].value == "email_001"
        assert workbook["Summary"]["A12"].value.startswith("=COUNTA")
        assert workbook["Summary"]["B4"].value.startswith("All cases;")

        now = datetime.now(timezone.utc)
        monthly_excel = client.get(
            f"/api/cases/export.xlsx?period=month&year={now.year}&month={now.month}"
        )
        assert monthly_excel.status_code == 200
        assert f"protozero-cases-{now:%Y-%m}.xlsx" in monthly_excel.headers["content-disposition"]
        monthly_workbook = load_workbook(
            BytesIO(monthly_excel.content), read_only=True, data_only=False
        )
        assert monthly_workbook["Summary"]["B4"].value.startswith(f"{now:%B %Y};")
        assert monthly_workbook["All Cases"].max_row <= 521

        annual_excel = client.get(f"/api/cases/export.xlsx?period=year&year={now.year}")
        assert annual_excel.status_code == 200
        annual_workbook = load_workbook(
            BytesIO(annual_excel.content), read_only=True, data_only=False
        )
        assert annual_workbook["Summary"]["B4"].value.startswith(f"Year {now.year};")

        invalid_month = client.get("/api/cases/export.xlsx?period=month&month=13")
        assert invalid_month.status_code == 422

        case = client.get("/api/cases/email_004")
        assert case.status_code == 200
        assert case.json()["comparisons"]
        assert case.json()["correlation_id"]

        evidence = client.get("/api/cases/email_004/document/SI")
        assert evidence.status_code == 200
        assert evidence.json()["text"]
        assert evidence.json()["attachment_path"].endswith("email_004_SI.txt")

        original = client.get("/api/cases/email_004/document/SI/raw")
        assert original.status_code == 200
        assert original.content

        events = client.get("/api/cases/email_004/events")
        assert events.status_code == 200
        assert events.json()["items"]
        assert all(
            event["correlation_id"] == case.json()["correlation_id"]
            for event in events.json()["items"]
        )
        assert events.json()["items"] == sorted(
            events.json()["items"], key=lambda event: event["seq"]
        )

        report = client.get("/api/cases/email_004/report.pdf")
        assert report.status_code == 200
        assert report.headers["content-type"] == "application/pdf"
        assert "attachment" in report.headers["content-disposition"]
        assert report.content.startswith(b"%PDF-")
        with fitz.open(stream=report.content, filetype="pdf") as document:
            report_text = "\n".join(page.get_text() for page in document)
        assert "email_004" in report_text
        assert "Field comparison" in report_text
        assert "Audit trail" in report_text

        decision = {
            "action": "reject",
            "label": "Reject the draft",
            "done": "Draft rejected.",
            "reviewer_id": "api-test",
        }
        saved = client.post("/api/cases/email_004/decision", json=decision)
        assert saved.status_code == 200
        assert saved.json()["decision"]["action"] == "reject"
        assert client.get("/api/cases/email_004/decision").json()["decision"] is not None

        removed = client.delete("/api/cases/email_004/decision?reviewer_id=api-test")
        assert removed.status_code == 200
        assert client.get("/api/cases/email_004/decision").json()["decision"] is None

        handoff = client.post(
            "/api/cases/email_004/decision",
            json={
                "action": "review",
                "label": "Send to a person",
                "done": "Sent to a person.",
                "reviewer_id": "api-test",
            },
        )
        assert handoff.status_code == 200
        queue = client.get("/api/review").json()["items"]
        assert any(item["id"] == "email_004:manual" for item in queue)
        client.delete("/api/cases/email_004/decision?reviewer_id=api-test")
        queue = client.get("/api/review").json()["items"]
        assert all(item["id"] != "email_004:manual" for item in queue)

        for trap_id in ["email_004", "email_506", "email_503", "email_300", "email_513"]:
            assert client.get(f"/api/cases/{trap_id}").status_code == 200


def test_assignment_and_review_status_are_audited_and_survive_rerun() -> None:
    with TestClient(app) as client:
        email_id = "email_010"
        assigned = client.patch(
            f"/api/cases/{email_id}/workflow",
            json={
                "assigned_to": "api-test",
                "review_status": "in_progress",
                "reviewer_id": "api-test",
            },
        )
        assert assigned.status_code == 200
        assert assigned.json()["case"]["assigned_to"] == "api-test"
        assert assigned.json()["case"]["review_status"] == "in_progress"

        rerun = client.post(f"/api/cases/{email_id}/rerun")
        assert rerun.status_code == 200
        assert rerun.json()["assigned_to"] == "api-test"
        assert rerun.json()["review_status"] == "in_progress"

        events = client.get(f"/api/cases/{email_id}/events").json()["items"]
        assert any(event["action"] == "CASE_ASSIGNED" for event in events)
        assert any(event["action"] == "REVIEW_STATUS_CHANGED" for event in events)

        cleared = client.patch(
            f"/api/cases/{email_id}/workflow",
            json={
                "assigned_to": None,
                "review_status": "unassigned",
                "reviewer_id": "api-test",
            },
        )
        assert cleared.status_code == 200


def test_correction_rejoins_at_comparison_without_rewriting_extraction() -> None:
    with TestClient(app) as client:
        client.post("/api/cases/email_516/rerun")
        before = client.get("/api/cases/email_516").json()
        comparison = next(
            value for value in before["comparisons"]
            if value["field"] == "gross_weight_kg"
        )
        original_document_value = next(
            document for document in before["documents"] if document["role"] == "BL"
        )["fields"]["gross_weight_kg"]["value"]

        resolved = client.post(
            "/api/review/email_516:gross_weight_kg/resolve",
            json={
                "action": "correct",
                "field": "gross_weight_kg",
                "document_role": "BL",
                "correct_value": comparison["si"]["value"],
                "reviewer_id": "api-test",
            },
        )
        assert resolved.status_code == 200
        after = resolved.json()["case"]
        corrected = next(
            value for value in after["comparisons"]
            if value["field"] == "gross_weight_kg"
        )
        assert corrected["bl"]["value"] == comparison["si"]["value"]
        assert corrected["bl"]["extracted_by"] == "human"
        assert corrected["human_reviewed"] is True
        assert next(
            document for document in after["documents"] if document["role"] == "BL"
        )["fields"]["gross_weight_kg"]["value"] == original_document_value

        events = client.get("/api/cases/email_516/events").json()["items"]
        assert any(event["action"] == "HUMAN_CORRECTION" for event in events)
        assert any(event["action"] == "COMPARISON_RERUN" for event in events)


def test_review_actions_update_queue_and_audit() -> None:
    with TestClient(app) as client:
        queue = client.get("/api/review").json()["items"]
        assert queue

        retry_item = next(
            item
            for item in queue
            if item["reason"] in {"UNREADABLE_DOCUMENT", "GROUNDING_FAILED", "PROCESSING_ERROR"}
        )
        retried = client.post(
            f"/api/review/{retry_item['id']}/retry", json={"force_llm": False}
        )
        assert retried.status_code == 200
        retry_events = client.get(
            f"/api/cases/{retry_item['email_id']}/events"
        ).json()["items"]
        assert any(event["action"] == "COMPARISON_RERUN" for event in retry_events)

        queue = client.get("/api/review").json()["items"]
        item = queue[0]

        resolved = client.post(
            f"/api/review/{item['id']}/resolve",
            json={
                "action": "confirm",
                "field": item["fields"][0] if item["fields"] else None,
                "correct_value": None,
                "reviewer_id": "api-test",
            },
        )
        assert resolved.status_code == 200
        assert resolved.json()["review_item"]["state"] == "resolved"

        events = client.get(f"/api/cases/{item['email_id']}/events").json()["items"]
        assert any(event["action"] == "HUMAN_CORRECTION" for event in events)
