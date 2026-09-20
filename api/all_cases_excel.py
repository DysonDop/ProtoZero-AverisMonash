"""Create one audit-friendly Excel workbook containing every case."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pipeline.schemas import AuditEvent, Case, CaseDecision, ReviewItem


INK = "101A24"
PAPER = "F5F2EC"
WHITE = "FFFFFF"
MUTED = "55606B"
RULE = "D9DCE0"
ORANGE = "D88C3D"
BLUE = "006DAE"
WRONG = "B83A28"
CLEAR = "1F6252"

FIELD_LABELS = {
    "shipper": "Shipper",
    "consignee": "Consignee",
    "notify_party": "Notify party",
    "port_of_loading": "Port of loading",
    "port_of_discharge": "Port of discharge",
    "container_count": "Container count",
    "gross_weight_kg": "Gross weight (KG)",
}


def build_all_cases_workbook(
    cases: list[Case],
    events_by_case: dict[str, list[AuditEvent]],
    reviews: list[ReviewItem],
    decisions: dict[str, CaseDecision],
    *,
    generated_at: datetime | None = None,
    scope_label: str = "All cases",
    date_range_label: str = "All available dates",
) -> bytes:
    """Return an XLSX with all structured case, evidence, review, and audit data."""
    generated_at = generated_at or datetime.now(timezone.utc)
    cases = sorted(cases, key=lambda case: case.email_id)

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "ProtoZero all-cases export"
    workbook.properties.subject = "Shipping document verification cases"
    workbook.properties.creator = "ProtoZero"

    all_cases = _build_all_cases_sheet(workbook, cases, decisions)
    _build_summary_sheet(
        workbook,
        cases,
        reviews,
        generated_at,
        all_cases.max_row,
        scope_label,
        date_range_label,
    )
    _build_comparisons_sheet(workbook, cases)
    _build_evidence_sheet(workbook, cases)
    _build_audit_sheet(workbook, cases, events_by_case)
    _build_reviews_sheet(workbook, reviews)

    workbook.active = workbook.sheetnames.index("Summary")
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _build_summary_sheet(
    workbook: Workbook,
    cases,
    reviews,
    generated_at,
    case_end_row,
    scope_label,
    date_range_label,
):
    sheet = workbook.create_sheet("Summary", 0)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F2")
    sheet["A1"] = "PROTOZERO - CASE REGISTER EXPORT"
    sheet["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=INK)
    sheet["A1"].alignment = Alignment(vertical="center")
    for row in sheet["A1:F2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=INK)

    sheet["A4"] = "Scope"
    sheet["B4"] = f"{scope_label}; worklist search, result filters and pagination do not limit this export."
    sheet["A5"] = "Report period"
    sheet["B5"] = (
        f"{date_range_label}; received timestamp, or first processing timestamp "
        "when the source email has no date."
    )
    sheet["A6"] = "Generated (UTC)"
    sheet["B6"] = _excel_datetime(generated_at)
    sheet["B6"].number_format = "yyyy-mm-dd hh:mm:ss"
    sheet["A7"] = "Included sheets"
    sheet["B7"] = "All Cases, Field Comparisons, Document Evidence, Audit Trail, Review Queue"
    sheet["A8"] = "Source files"
    sheet["B8"] = "Attachment paths and extracted evidence are included; original file binaries remain in the application."

    for cell in sheet["A4:A8"]:
        cell[0].font = Font(name="Aptos", bold=True, color=MUTED)
    for row in sheet["B4:B8"]:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    sheet["A10"] = "OPERATIONAL TOTALS"
    sheet["A10"].font = Font(name="Aptos", size=11, bold=True, color=MUTED)
    formula_end_row = max(case_end_row, 2)
    metrics = [
        ("Cases", f"=COUNTA('All Cases'!A2:A{formula_end_row})"),
        ("Comparison cases", f'=COUNTIF(\'All Cases\'!F2:F{formula_end_row},"BL_COMPARISON")'),
        ("All clear", f'=COUNTIF(\'All Cases\'!H2:H{formula_end_row},"OK")'),
        ("Differences found", f'=COUNTIF(\'All Cases\'!H2:H{formula_end_row},"MISMATCH")'),
        ("Needs a person", f'=COUNTIF(\'All Cases\'!H2:H{formula_end_row},"NEEDS_REVIEW")'),
        ("Open review items", str(sum(1 for review in reviews if review.state == "open"))),
    ]
    for index, (label, value) in enumerate(metrics):
        column = 1 + (index % 3) * 2
        row = 11 + (index // 3) * 3
        label_cell = sheet.cell(row=row, column=column, value=label)
        value_cell = sheet.cell(row=row + 1, column=column, value=value)
        sheet.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + 1)
        sheet.merge_cells(start_row=row + 1, start_column=column, end_row=row + 1, end_column=column + 1)
        for block_row in (row, row + 1):
            for block_column in (column, column + 1):
                sheet.cell(block_row, block_column).fill = PatternFill("solid", fgColor=PAPER)
        label_cell.font = Font(name="Aptos", size=9, bold=True, color=MUTED)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.font = Font(name="Aptos Display", size=18, bold=True, color=INK)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = "#,##0"

    sheet["A18"] = "How to use this workbook"
    sheet["A18"].font = Font(name="Aptos Display", size=13, bold=True, color=INK)
    instructions = [
        "Start in All Cases to filter by result, lifecycle, sender, or case ID.",
        "Use Field Comparisons for SI-versus-BL values, similarity, confidence, and outcomes.",
        "Use Document Evidence for extraction method, source text, and exact page or line locations.",
        "Use Audit Trail for correlation IDs, timestamps, actors, changes, and reasons.",
        "Use Review Queue to see open and resolved human-review items.",
    ]
    for row, instruction in enumerate(instructions, start=19):
        sheet.cell(row=row, column=1, value=f"{row - 18}.")
        sheet.cell(row=row, column=2, value=instruction)
        sheet.cell(row=row, column=2).alignment = Alignment(wrap_text=True)

    widths = {"A": 20, "B": 28, "C": 6, "D": 20, "E": 28, "F": 6}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 24
    sheet.freeze_panes = "A4"


def _build_all_cases_sheet(workbook: Workbook, cases: list[Case], decisions):
    headers = [
        "Case ID", "Correlation ID", "Received (UTC)", "Sender", "Subject", "Category",
        "Category confidence", "Result", "Has defect", "Defect fields", "Escalation reasons",
        "Lifecycle", "Summary", "Reviewer decision", "Decision label", "Decision reviewer",
        "Decision recorded (UTC)", "Documents", "Comparisons", "Pipeline version",
        "Prompt version", "Models", "Timings (ms)", "Created (UTC)", "Updated (UTC)",
    ]
    rows = []
    for case in cases:
        decision = decisions.get(case.email_id)
        rows.append([
            case.email_id, case.correlation_id, _excel_datetime(case.received_at), case.from_addr,
            case.subject, case.category, case.category_confidence, case.status, case.has_defect,
            ", ".join(case.defect_fields), ", ".join(case.escalation_reasons), case.lifecycle,
            case.summary, decision.action if decision else None, decision.label if decision else None,
            decision.reviewer_id if decision else None,
            _excel_datetime(decision.recorded_at) if decision else None,
            len(case.documents), len(case.comparisons), case.pipeline_version, case.prompt_version,
            json.dumps(case.models, sort_keys=True), json.dumps(case.timings_ms, sort_keys=True),
            _excel_datetime(case.created_at), _excel_datetime(case.updated_at),
        ])
    sheet = _data_sheet(workbook, "All Cases", headers, rows, freeze="A2", table_name="AllCasesTable")
    _set_widths(sheet, [15, 38, 20, 28, 52, 22, 20, 20, 12, 30, 32, 15, 64, 18, 26, 22, 22, 12, 14, 18, 16, 30, 30, 20, 20])
    for row in range(2, sheet.max_row + 1):
        for column in (3, 17, 24, 25):
            sheet.cell(row, column).number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(row, 7).number_format = "0%"
    if sheet.max_row > 1:
        _status_rules(sheet, f"H2:H{sheet.max_row}")
    return sheet


def _build_comparisons_sheet(workbook: Workbook, cases: list[Case]):
    headers = [
        "Case ID", "Field", "Instruction value", "Draft value", "Verdict", "Similarity",
        "Machine confidence", "Hard fail", "Human reviewed", "Explanation",
        "Instruction normalized", "Draft normalized", "Instruction evidence location",
        "Draft evidence location", "Instruction evidence", "Draft evidence",
    ]
    rows = []
    for case in cases:
        for comparison in case.comparisons:
            rows.append([
                case.email_id, FIELD_LABELS.get(comparison.field, comparison.field),
                comparison.si.value, comparison.bl.value, comparison.verdict,
                comparison.similarity / 100 if comparison.similarity is not None else None,
                comparison.confidence.score, comparison.confidence.hard_fail,
                comparison.human_reviewed, comparison.explanation,
                comparison.si_normalized, comparison.bl_normalized,
                _locator(comparison.si.locator), _locator(comparison.bl.locator),
                comparison.si.evidence, comparison.bl.evidence,
            ])
    sheet = _data_sheet(workbook, "Field Comparisons", headers, rows, freeze="C2", table_name="ComparisonsTable")
    _set_widths(sheet, [15, 24, 34, 34, 18, 14, 20, 24, 16, 55, 30, 30, 28, 28, 55, 55])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 6).number_format = "0%"
        sheet.cell(row, 7).number_format = "0%"
    if sheet.max_row > 1:
        _status_rules(sheet, f"E2:E{sheet.max_row}")


def _build_evidence_sheet(workbook: Workbook, cases: list[Case]):
    headers = [
        "Case ID", "Document role", "Attachment", "Detected type", "Format", "Readable",
        "Page count", "Field", "Value", "Label seen", "Extracted by", "Service confidence",
        "Model confidence", "Evidence", "Page", "Sheet", "Line", "Character start",
        "Character end", "Bounding box", "Text SHA-256", "Parse error",
    ]
    rows = []
    for case in cases:
        for document in case.documents:
            if document.fields:
                for field_name, field in document.fields:
                    locator = field.locator
                    rows.append([
                        case.email_id, document.role, Path(document.attachment_path).name,
                        document.detected_kind, document.fmt, document.readable, document.page_count,
                        FIELD_LABELS.get(field_name, field_name), field.value, field.label_seen,
                        field.extracted_by, field.service_confidence, field.model_confidence,
                        field.evidence, locator.page if locator else None, locator.sheet if locator else None,
                        (locator.line + 1) if locator and locator.line is not None else None,
                        locator.char_start if locator else None, locator.char_end if locator else None,
                        json.dumps(locator.bbox) if locator and locator.bbox else None,
                        document.text_sha256, document.parse_error,
                    ])
            else:
                rows.append([
                    case.email_id, document.role, Path(document.attachment_path).name,
                    document.detected_kind, document.fmt, document.readable, document.page_count,
                ] + [None] * 15)
    sheet = _data_sheet(workbook, "Document Evidence", headers, rows, freeze="H2", table_name="EvidenceTable")
    _set_widths(sheet, [15, 16, 34, 22, 12, 12, 12, 24, 34, 24, 20, 20, 18, 60, 10, 18, 10, 16, 16, 24, 66, 40])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 12).number_format = "0%"
        sheet.cell(row, 13).number_format = "0%"


def _build_audit_sheet(workbook: Workbook, cases, events_by_case):
    headers = [
        "Case ID", "Correlation ID", "Sequence", "Timestamp (UTC)", "Actor", "Reviewer ID",
        "Action", "Field", "Previous value", "New value", "Reason",
    ]
    rows = []
    for case in cases:
        for event in events_by_case.get(case.email_id, []):
            rows.append([
                event.email_id, event.correlation_id, event.seq, _excel_datetime(event.at),
                event.actor, event.reviewer_id, event.action,
                FIELD_LABELS.get(event.field, event.field) if event.field else None,
                event.previous_value, event.new_value, event.reason,
            ])
    sheet = _data_sheet(workbook, "Audit Trail", headers, rows, freeze="C2", table_name="AuditTable")
    _set_widths(sheet, [15, 38, 12, 20, 14, 22, 30, 24, 34, 34, 70])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 4).number_format = "yyyy-mm-dd hh:mm:ss"


def _build_reviews_sheet(workbook: Workbook, reviews: list[ReviewItem]):
    headers = [
        "Review ID", "Case ID", "State", "Reason", "Fields", "Reason detail",
        "Instruction value", "Draft value", "Confidence", "Instruction evidence",
        "Draft evidence", "Created (UTC)", "Resolved (UTC)", "Resolved by",
    ]
    rows = [[
        review.id, review.email_id, review.state, review.reason, ", ".join(review.fields),
        review.reason_detail, review.si_value, review.bl_value, review.confidence,
        review.si_evidence, review.bl_evidence, _excel_datetime(review.created_at),
        _excel_datetime(review.resolved_at), review.resolved_by,
    ] for review in sorted(reviews, key=lambda item: (item.state, item.email_id, item.id))]
    sheet = _data_sheet(workbook, "Review Queue", headers, rows, freeze="C2", table_name="ReviewTable")
    _set_widths(sheet, [28, 15, 14, 28, 30, 65, 34, 34, 15, 55, 55, 20, 20, 22])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 9).number_format = "0%"
        sheet.cell(row, 12).number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(row, 13).number_format = "yyyy-mm-dd hh:mm:ss"


def _data_sheet(workbook, name, headers, rows, *, freeze, table_name):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_cell(value) for value in row])
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    header = sheet[1]
    for cell in header:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    stripe = PatternFill("solid", fgColor="F7F9FA")
    for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_number % 2 == 0:
                cell.fill = stripe
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def _set_widths(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 70)


def _status_rules(sheet, cell_range):
    first = cell_range.split(":")[0]
    thin = Side(style="thin", color=RULE)
    rules = [
        ("MISMATCH", "FFF0ED", WRONG),
        ("NEEDS_REVIEW", "EEF6FB", BLUE),
        ("REVIEW", "EEF6FB", BLUE),
        ("OK", "EFF7F4", CLEAR),
        ("MATCH", "EFF7F4", CLEAR),
    ]
    for value, fill, font in rules:
        sheet.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[f'{first}="{value}"'],
                fill=PatternFill("solid", fgColor=fill),
                font=Font(color=font, bold=True),
                border=Border(bottom=thin),
            ),
        )


def _locator(locator) -> str | None:
    if locator is None:
        return None
    if locator.page is not None:
        return f"page {locator.page}"
    if locator.sheet:
        return f"{locator.sheet}, row {(locator.line + 1) if locator.line is not None else '-'}"
    if locator.line is not None:
        return f"line {locator.line + 1}"
    return None


def _excel_datetime(value):
    if value is None:
        return None
    if value.tzinfo is not None:
        value = value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def _safe_cell(value):
    if not isinstance(value, str):
        return value
    value = value.replace("\x00", "")[:32767]
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
